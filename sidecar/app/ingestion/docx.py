"""DOCX extraction.

Text comes from python-docx; images come from the underlying ZIP. python-docx
exposes ``document.inline_shapes`` but resolving those to bytes means walking
relationship parts anyway, and it misses floating (anchored) images entirely --
reading ``word/media/`` gets everything in one pass.

DOCX has no pages: pagination is computed by the renderer, not stored in the
file. So ``page_number`` stays null for DOCX chunks, and the UI opens the
document rather than a page.
"""

from __future__ import annotations

import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.pdf import MIN_IMAGE_PIXELS, ExtractedImage
from app.ingestion.titles import clean_title
from app.logging_conf import get_logger
from app.storage import files as storage

log = get_logger(__name__)

_MEDIA_PREFIX = "word/media/"
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}


@dataclass
class DocxExtraction:
    text: str = ""
    images: list[ExtractedImage] = field(default_factory=list)
    title: str | None = None


def _iter_block_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in document order.

    python-docx exposes ``.paragraphs`` and ``.tables`` as separate flat lists,
    which loses their interleaving -- a table's contents would end up detached
    from the text introducing it. Walking the body XML preserves the order.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _table_to_text(table: Table) -> str:
    """Flatten a table to one line per row.

    Tab-separated cells keep the row grouped as a unit for embedding, which is
    usually what a table row means, while staying readable in a search snippet.
    """
    lines = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        # Merged cells repeat their text across the span; collapse the repeats.
        deduped: list[str] = []
        for cell in cells:
            if not deduped or cell != deduped[-1]:
                deduped.append(cell)
        line = "\t".join(c for c in deduped if c)
        if line:
            lines.append(line)
    return "\n".join(lines)


def _extract_images(path: Path, sha256: str) -> list[ExtractedImage]:
    from PIL import Image, UnidentifiedImageError

    images: list[ExtractedImage] = []
    with zipfile.ZipFile(path) as archive:
        media = sorted(
            name
            for name in archive.namelist()
            if name.startswith(_MEDIA_PREFIX) and Path(name).suffix.lower() in _IMAGE_SUFFIXES
        )
        for ordinal, name in enumerate(media):
            data = archive.read(name)
            destination = storage.derived_path(
                sha256, "images", f"img-{ordinal:04d}{Path(name).suffix.lower()}"
            )
            destination.write_bytes(data)

            try:
                with Image.open(destination) as image:
                    width, height = image.size
            except (UnidentifiedImageError, OSError) as exc:
                log.debug("skipping unreadable docx image %s: %s", name, exc)
                destination.unlink(missing_ok=True)
                continue

            if width < MIN_IMAGE_PIXELS or height < MIN_IMAGE_PIXELS:
                destination.unlink(missing_ok=True)
                continue

            images.append(
                ExtractedImage(
                    path=destination,
                    page_number=None,
                    ordinal=ordinal,
                    width=width,
                    height=height,
                )
            )
    return images


def extract(path: Path, sha256: str) -> DocxExtraction:
    try:
        document = docx.Document(str(path))
    except (KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ValueError(f"This DOCX file could not be read: {exc}") from exc

    parts: list[str] = []
    for block in _iter_block_items(document):
        if isinstance(block, Table):
            table_text = _table_to_text(block)
            if table_text:
                parts.append(table_text)
        else:
            text = block.text.strip()
            if text:
                parts.append(text)

    result = DocxExtraction(
        text="\n\n".join(parts),
        images=_extract_images(path, sha256),
        title=clean_title(document.core_properties.title, path.name),
    )
    log.info("docx %s: %d chars, %d images", path.name, len(result.text), len(result.images))
    return result
